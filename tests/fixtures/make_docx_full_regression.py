"""合成一份"全容器 + 假敏感数据"的 docx —— 供独立验证做真实场景回归。

背景
----
P1~P6 的样例是**按容器拆开**的（一份文档只测一类结构），
独立验证拿真实文档跑时遇到两个问题：

1. 真实文档早已把敏感信息替换过-> 0 命中，验证不到任何东西；
2. 真实文档不含页眉页脚/ 批注 / 文本框 / 脚注尾注 / 校对标记
   -> 覆盖不到 P4/P5/P6 的能力。

本脚本合成**一份**文档，同时覆盖全部容器，且刻意做成"一份商务文档"
（投标响应文件），而不是把测试用例堆在一起——这样容器之间的组合关系
（表格里套表格、文本框嵌在有批注的段落里……）也能被一起验证。

容器清单
--------
====================  ====================================================
容器                  覆盖内容
====================  ====================================================
正文段落              跨 run、跨 run+超链接边界、``<w:br/>``、``<w:tab/>``
超链接                mailto（Target 需改写）、https、显示文字含邮箱、
                      内部锚点
表格                  普通表格、单元格多段落、合并单元格、嵌套表格
页眉页脚              3 节（第3节链接到前一节）、首页不同、偶数页不同、
                      页眉内表格、页眉内 mailto 链接
批注                  含敏感 / 不含敏感 / 多段落三种
文本框                外层段落同时含普通文本与文本框、**嵌套文本框**、
                      文本框内表格
脚注 / 尾注                正文带引用标记；part 内含系统保留条目
校对标记                ``<w:proofErr>``、``<w:bookmarkStart/End>``
域代码（对照组）        已知限制：应被跳过并计入报告
====================  ====================================================

数据来源
--------
**全部复用** :mod:`fixtures.make_samples` 里既有的 ``FAKE_*`` 常量，
不引入任何新的假数据，也不含任何真实个人信息。

运行方式::

    python tests/fixtures/make_docx_full_regression.py
    # 生成 samples/docx_demo_full_regression.docx，并打印容器/敏感值对照表

注意：本脚本只写``samples/``，不碰 ``src/``。
"""

from __future__ import annotations

import copy
import sys
from pathlib import Path
from typing import Dict, List, Sequence, Tuple

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

_TESTS_DIR = Path(__file__).resolve().parents[1]
if str(_TESTS_DIR) not in sys.path:
    sys.path.insert(0, str(_TESTS_DIR))

from fixtures.make_docx_samples import (  # noqa: E402
    _attach_note_ref,
    _attach_notes,
    _make_textbox,
    add_hyperlink,
)
from fixtures.make_samples import (  # noqa: E402
    FAKE_BANK_CARD,
    FAKE_EMAIL,
    FAKE_ID_VALID,
    FAKE_IP,
    FAKE_PHONE_1,
    FAKE_PHONE_2,
)

PROJECT_ROOT = Path(__file__).resolve().parents[2]
SAMPLES_DIR = PROJECT_ROOT / "samples"
OUTPUT_NAME = "docx_demo_full_regression.docx"

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

#:容器 -> 该容器内出现的敏感值（供独立验证逐条对照）
EXPECTED_SENSITIVE: Dict[str, List[str]] = {}


def _note(container: str, *values: str) -> None:
    """登记"某容器内应出现这些敏感值"，最后打印成对照表。"""
    bucket = EXPECTED_SENSITIVE.setdefault(container, [])
    for value in values:
        if value not in bucket:
            bucket.append(value)


# ---------------------------------------------------------------------------
# 低层构造helper
# ---------------------------------------------------------------------------
def _split_runs(paragraph, text: str, cuts: Sequence[int], styles: Sequence[str]) -> None:
    """把 ``text`` 按 ``cuts`` 切成多个 run，并分别施加格式。

    刻意在**号码中间**下刀，确保任何单run 匹配都不可能命中——
    这才是跨 run 处理真正要解决的场景。
    """
    marks = [0, *cuts, len(text)]
    for index in range(len(marks) - 1):
        chunk = text[marks[index] : marks[index + 1]]
        if not chunk:
            continue
        run = paragraph.add_run(chunk)
        style = styles[index % len(styles)]
        if style == "bold":
            run.bold = True
        elif style == "italic":
            run.italic = True
        elif style == "underline":
            run.underline = True


def _wrap_prooferr(paragraph_element, text: str, err_type: str = "spellStart") -> None:
    """插入一对 ``<w:proofErr>``，中间夹一个只含 ``text`` 的 run。

    Word 默认会自动插入校对标记，几乎所有人工编辑过的文档都有；
    这类标记是**空元素、不占偏移**，必须放行否则整段会被跳过。
    """
    start = OxmlElement("w:proofErr")
    start.set(qn("w:type"), f"{err_type}Start")
    paragraph_element.append(start)

    run = OxmlElement("w:r")
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    paragraph_element.append(run)

    end = OxmlElement("w:proofErr")
    end.set(qn("w:type"), f"{err_type}End")
    paragraph_element.append(end)


def _wrap_bookmark(
    paragraph_element,
    text: str,
    bookmark_id: str,
    bookmark_name: str,
) -> None:
    """插入一对 ``<w:bookmarkStart>`` / ``<w:bookmarkEnd>``，中间夹文本。

    ``_GoBack`` 是 Word 的隐藏系统书签，真实文档里几乎必然存在。
    """
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), bookmark_id)
    start.set(qn("w:name"), bookmark_name)
    paragraph_element.append(start)

    run = OxmlElement("w:r")
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    paragraph_element.append(run)

    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), bookmark_id)
    paragraph_element.append(end)


def _attach_comments_at(document, anchors) -> None:
    """给**指定段落**挂批注（python-docx 无此 API，手工构造）。

    与 ``make_docx_samples.attach_comments`` 的差别：那个函数会自己新建一个段落，
    这里必须把锚点插到已有段落上，才能验证"带批注的正文段落不被整段跳过"。

    Args:
        document: 目标 Document。
        anchors: ``[(段落, [(批注id, 作者, initials, [段落文本, ...]), ...]), ...]``。
            同一段落可挂多条批注；同一批注可有多个段落。
    """
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part

    entries: List[str] = []
    for paragraph, comments in anchors:
        for comment_id, author, initials, texts in comments:
            body = "".join(
                f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p>" for text in texts
            )
            entries.append(
                f'<w:comment w:id="{comment_id}" w:author="{author}" '
                f'w:initials="{initials}" w:date="2026-01-01T00:00:00Z">'
                f"{body}</w:comment>"
            )

            start = OxmlElement("w:commentRangeStart")
            start.set(qn("w:id"), str(comment_id))
            paragraph._p.insert(0, start)

            end = OxmlElement("w:commentRangeEnd")
            end.set(qn("w:id"), str(comment_id))
            paragraph._p.append(end)

            run = OxmlElement("w:r")
            rpr = OxmlElement("w:rPr")
            style = OxmlElement("w:rStyle")
            style.set(qn("w:val"), "CommentReference")
            rpr.append(style)
            run.append(rpr)
            ref = OxmlElement("w:commentReference")
            ref.set(qn("w:id"), str(comment_id))
            run.append(ref)
            paragraph._p.append(run)

    if not entries:
        return

    xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w:comments xmlns:w="{W_NS}">' + "".join(entries) + "</w:comments>"
    )
    part = Part(
        PackURI("/word/comments.xml"),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml",
        xml.encode("utf-8"),
        document.part.package,
    )
    document.part.relate_to(part, RT.COMMENTS)


def _detached_table_element(document, rows: int, cols: int, texts: Sequence[str]):
    """造一张真表格再把 XML 摘出来，用于塞进文本框。

    直接手工拼 ``<w:tbl>`` 容易漏 ``tblPr`` / ``tblGrid`` 导致 Word 打不开；
    这里复用 python-docx 生成的完整结构，深拷贝后立刻从正文摘掉。
    """
    table = document.add_table(rows=rows, cols=cols)
    for index, cell in enumerate(_iter_cells(table)):
        cell.text = texts[index]
    element = copy.deepcopy(table._tbl)
    table._tbl.getparent().remove(table._tbl)
    return element


def _iter_cells(table):
    """按行展开单元格（合并格会重复出现，此处不使用合并）。"""
    for row in table.rows:
        for cell in row.cells:
            yield cell


def _make_field_paragraph(paragraph, label: str, instruction: str = " PAGE ") -> None:
    """插入一个域代码段落——**已知限制对照组**，应被跳过并计入报告。"""
    paragraph.add_run(label)
    begin_run = OxmlElement("w:r")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin_run.append(begin)
    paragraph._p.append(begin_run)

    instr_run = OxmlElement("w:r")
    instr = OxmlElement("w:instrText")
    instr.text = instruction
    instr_run.append(instr)
    paragraph._p.append(instr_run)

    sep_run = OxmlElement("w:r")
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    sep_run.append(sep)
    paragraph._p.append(sep_run)

    value_run = OxmlElement("w:r")
    value = OxmlElement("w:t")
    value.text = "1"
    value_run.append(value)
    paragraph._p.append(value_run)

    end_run = OxmlElement("w:r")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run.append(end)
    paragraph._p.append(end_run)


# ---------------------------------------------------------------------------
# 文档组装
# ---------------------------------------------------------------------------
def build() -> Path:
    """生成全容器回归样例并返回路径。"""
    SAMPLES_DIR.mkdir(parents=True, exist_ok=True)
    path = SAMPLES_DIR / OUTPUT_NAME

    document = Document()

    _build_cover(document)
    _build_body(document)
    _build_tech_section(document)
    _build_textboxes(document)
    _build_notes(document)
    _build_header_footer(document)

    document.save(str(path))
    return path


def _build_cover(document) -> None:
    """封面：标题 + 项目编号 + 联系方式（跨 run、br/tab、超链接）。"""
    title = document.add_heading("城市供水设施智能化改造工程", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("投标响应文件（商务分）")
    run.bold = True

    meta = document.add_paragraph()
    meta.add_run("项目编号：CQ-2026-SB-0147　　编制单位：示例市政工程有限公司")
    meta.add_run().add_tab()
    meta.add_run("编制日期：2026 年 10 月")

    # 跨 run 联系方式：手机号拆成 3 段（加粗/ 斜体 / 常规）
    contact = document.add_paragraph()
    contact.add_run("项目联系人：周宁　　联系电话：")
    _split_runs(contact, FAKE_PHONE_1, cuts=(3, 7), styles=("bold", "italic", "regular"))
    contact.add_run("　　备用号码：")
    contact.add_run(FAKE_PHONE_2)
    _note("正文段落", FAKE_PHONE_1, FAKE_PHONE_2)

    # 邮箱跨 run（在 @ 处断开）
    mail = document.add_paragraph()
    mail.add_run("电子邮箱：")
    _split_runs(
        mail,
        FAKE_EMAIL,
        cuts=(FAKE_EMAIL.index("@"),),
        styles=("bold", "underline"),
    )
    _note("正文段落", FAKE_EMAIL)

    # br / tab 段落
    address = document.add_paragraph()
    address.add_run("通讯地址：重庆市渝中区示例路 1 号")
    address.add_run().add_break(WD_BREAK.LINE)
    address.add_run("技术支持专线：")
    address.add_run(FAKE_IP)
    address.add_run().add_tab()
    address.add_run("（工作日 9:00-18:00）")
    _note("正文段落（br/tab）", FAKE_IP)

    # 超链接四种情形
    link_line = document.add_paragraph()
    link_line.add_run("在线沟通：")
    add_hyperlink(link_line, f"mailto:{FAKE_EMAIL}", "发送邮件给我们")
    link_line.add_run("　｜　")
    add_hyperlink(link_line, "https://example.com/portal", "供应商门户")
    link_line.add_run("　｜　")
    add_hyperlink(link_line, "", "跳转至附件清单", anchor="appendix")
    _note("超链接（mailto 目标）", FAKE_EMAIL)

    # 显示文字本身含邮箱 + 手机号跨越"普通 run / 链接文字"边界
    boundary = document.add_paragraph()
    boundary.add_run("备用联系邮箱 ")
    add_hyperlink(boundary, "https://example.com/mail", FAKE_EMAIL)
    boundary.add_run("，值班电话 1380013")
    add_hyperlink(boundary, "https://example.com/hotline", "8000（转分机）")
    _note("超链接（显示文字/跨边界）", FAKE_EMAIL)


def _build_body(document) -> None:
    """正文主体：商务条款 + 表格（普通 / 合并 / 嵌套）。"""
    document.add_heading("一、商务响应说明", level=1)

    document.add_paragraph(
        "我方已完整阅读招标文件全部内容，承诺以下条款："
    )
    numbered = document.add_paragraph("1. 投标有效期为 90 日历天。")
    numbered.paragraph_format.left_indent = None
    document.add_paragraph("2. 质保期为验收合格之日起 24 个月。")
    document.add_paragraph("3. 逾期送达的投标文件一律不予受理。")

    # ---- 表格1：普通表格 + 单元格多段落 ----
    document.add_heading("1.1 供应商联络信息表", level=2)
    table = document.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    table.cell(0, 0).text = "法定代表人"
    table.cell(0, 1).text = "周宁"
    table.cell(0, 2).text = f"身份证 {FAKE_ID_VALID}"
    table.cell(1, 0).text = "项目负责人"
    table.cell(1, 1).text = "陈晓"
    table.cell(1, 2).text = f"手机 {FAKE_PHONE_1}"
    table.cell(2, 0).text = "财务对接人"
    table.cell(2, 1).text = "李文静"
    table.cell(2, 2).text = f"邮箱 {FAKE_EMAIL}"
    # 单元格内多段落
    table.cell(2, 2).add_paragraph(f"备用手机 {FAKE_PHONE_2}")
    _note("表格（普通）", FAKE_ID_VALID, FAKE_PHONE_1, FAKE_EMAIL, FAKE_PHONE_2)

    # ---- 表格2：合并单元格 ----
    document.add_heading("1.2 开户信息", level=2)
    merged = document.add_table(rows=3, cols=2)
    merged.style = "Table Grid"
    merged.cell(0, 0).merge(merged.cell(0, 1))
    merged.cell(0, 0).text = "本企业对公账户信息（合并表头）"
    merged.cell(1, 0).text = "开户银行"
    merged.cell(1, 1).text = "示例银行重庆渝中支行"
    merged.cell(2, 0).text = "账号"
    merged.cell(2, 1).text = f"银行卡 {FAKE_BANK_CARD}"
    _note("表格（合并单元格）", FAKE_BANK_CARD)

    # ---- 表格3：嵌套表格 ----
    document.add_heading("1.3 现场服务安排", level=2)
    nested = document.add_table(rows=1, cols=1)
    nested.style = "Table Grid"
    outer_cell = nested.cell(0, 0)
    outer_cell.text = "服务组人员配置"
    inner = outer_cell.add_table(rows=2, cols=2)
    inner.style = "Table Grid"
    inner.cell(0, 0).text = "值守工程师"
    inner.cell(0, 1).text = f"电话 {FAKE_PHONE_2}"
    inner.cell(1, 0).text = "中控室"
    inner.cell(1, 1).text = f"内网地址 {FAKE_IP}"
    _note("表格（嵌套表）", FAKE_PHONE_2, FAKE_IP)

    # ---- 表格4：单元格内超链接（mailto 目标需改写） ----
    document.add_heading("1.4 资料接收", level=2)
    linked = document.add_table(rows=1, cols=1)
    linked.style = "Table Grid"
    add_hyperlink(
        linked.cell(0, 0).paragraphs[0],
        f"mailto:{FAKE_EMAIL}",
        "将投标资料发送至我方邮箱",
    )
    _note("表格（单元格内超链接）", FAKE_EMAIL)

    # ---- 域代码对照组（已知限制，应被跳过） ----
    document.add_paragraph()
    field_paragraph = document.add_paragraph()
    _make_field_paragraph(field_paragraph, "本节共 4 页，第 ")
    document.add_paragraph("（上句为域代码，属已知限制，预期被跳过并计入报告）")

    # ---- 批注挂载：挂在有真实上文的两段上 ----
    document.add_heading("二、澄清与承诺", level=1)
    annotated = document.add_paragraph(
        "我方承诺在收到中标通知书后 5 个工作日内完成合同签订，"
        f"联系人电话 {FAKE_PHONE_1} 保持畅通。"
    )
    _note("正文段落（带批注）", FAKE_PHONE_1)

    clean_annotated = document.add_paragraph("质保期内响应时间不超过 2 小时。")

    _attach_comments_at(
        document,
        [
            (
                annotated,
                [
                    (1, "周宁", "ZN", [f"请核对手机号 {FAKE_PHONE_1} 是否为最新号码。"]),
                    (2, "陈晓", "CX", ["此条无需修改。"]),
                    (3, "李文静", "LWJ", ["签署页引用。", f"开户账号 {FAKE_BANK_CARD} 请再次核对。"]),
                ],
            ),
            (clean_annotated, [(4, "周宁", "ZN", ["响应时间承诺偏保守，可保留。"])]),
        ],
    )
    _note("批注", FAKE_PHONE_1, FAKE_BANK_CARD)


def _build_tech_section(document) -> None:
    """技术章节：校对标记 + 书签包裹的段落。"""
    document.add_heading("三、技术方案要点", level=1)

    para = document.add_paragraph()
    _wrap_prooferr(
        para._p,
        f"本项目拟部署 12 处智能监测点，负责人电话 {FAKE_PHONE_1}。",
    )
    _note("正文段落（proofErr 包裹）", FAKE_PHONE_1)

    para2 = document.add_paragraph()
    _wrap_prooferr(
        para2._p,
        f"数据上报网关地址 {FAKE_IP}，端口 8080。",
        err_type="gramStart",
    )
    _note("正文段落（proofErr 包裹）", FAKE_IP)

    para3 = document.add_paragraph()
    _wrap_bookmark(
        para3._p,
        f"备品备件仓库设于厂区东侧，押运负责人邮箱 {FAKE_EMAIL}。",
        bookmark_id="0",
        bookmark_name="_GoBack",
    )
    _note("正文段落（_GoBack 书签）", FAKE_EMAIL)

    para4 = document.add_paragraph()
    _wrap_bookmark(
        para4._p,
        f"合同条款详见附件 A，签署人身份证 {FAKE_ID_VALID}。",
        bookmark_id="1",
        bookmark_name="appendix",
    )
    _note("正文段落（普通书签）", FAKE_ID_VALID)

    document.add_paragraph(
        "本章技术方案不响应招标文件第 4.2 条，已在偏离表中列明。"
    )


def _build_textboxes(document) -> None:
    """文本框：外层混合、嵌套、框内表格。"""
    document.add_heading("四、现场提示", level=1)

    # 外层段落同时含普通文本与文本框——两者都要脱敏
    outer = document.add_paragraph()
    outer.add_run(f"进入厂区前请拨打值班电话 {FAKE_PHONE_2}，")
    outer._p.append(_make_textbox(f"文本框：应急联络邮箱 {FAKE_EMAIL}"))
    _note("正文段落", FAKE_PHONE_2)
    _note("文本框 1", FAKE_EMAIL)

    # 嵌套文本框：外层文本框里再放一个文本框
    nested_holder = document.add_paragraph()
    outer_box = _make_textbox("外层提示框")
    inner_para = OxmlElement("w:p")
    inner_box = _make_textbox(f"内层提示框：值班邮箱 {FAKE_EMAIL}")
    inner_para.append(inner_box)
    outer_box.append(inner_para)
    nested_holder._p.append(outer_box)
    _note("文本框 3（嵌套）", FAKE_EMAIL)

    # 文本框内表格
    box_holder = document.add_paragraph()
    box_with_table = _make_textbox("附件清单概览")
    table_element = _detached_table_element(
        document,
        rows=2,
        cols=2,
        texts=[
            "附件名称",
            "份数",
            f"报价明细（卡号 {FAKE_BANK_CARD}）",
            "1",
        ],
    )
    box_with_table.append(table_element)
    box_holder._p.append(box_with_table)
    _note("文本框内表格", FAKE_BANK_CARD)


def _build_notes(document) -> None:
    """脚注与尾注：正文带引用标记，part 内含系统保留条目。"""
    document.add_heading("五、附注", level=1)

    note_paragraph = document.add_paragraph(
        "本项目已通过初步评审，核查联系人电话 "
    )
    note_paragraph.add_run(FAKE_PHONE_1)
    _attach_note_ref(note_paragraph, "footnoteReference", "2")
    note_paragraph.add_run("；归档联系人邮箱 ")
    note_paragraph.add_run(FAKE_EMAIL)
    _attach_note_ref(note_paragraph, "endnoteReference", "2")
    _note("正文段落（脚注/尾注引用）", FAKE_PHONE_1, FAKE_EMAIL)

    _attach_notes(
        document,
        footnotes=[
            (2, f"核查记录见附件 B，联系人手机 {FAKE_PHONE_2}。"),
            (3, "本脚注不含敏感信息。"),
        ],
        endnotes=[
            (2, f"归档服务器 {FAKE_IP}，由运维组统一保管。"),
            (3, f"签署人身份证 {FAKE_ID_VALID} 已归档。"),
        ],
    )
    _note("脚注", FAKE_PHONE_2)
    _note("尾注", FAKE_IP, FAKE_ID_VALID)

    document.add_paragraph("（本页无正文内容，用于确认空段落不被误报。）")


def _build_header_footer(document) -> None:
    """页眉页脚：多节 + 首页不同 + 偶数页不同 + 页眉表格 + 页眉链接。"""
    section1 = document.sections[0]

    section1.header.is_linked_to_previous = False
    section1.header.paragraphs[0].text = f"投标响应文件（商务分）　联系人电话 {FAKE_PHONE_1}"
    section1.footer.is_linked_to_previous = False
    section1.footer.paragraphs[0].text = f"编制单位：示例市政工程有限公司　邮箱 {FAKE_EMAIL}"
    _note("页眉（第1节）", FAKE_PHONE_1)
    _note("页脚（第1节）", FAKE_EMAIL)

    # 页脚里的域代码（真实文档的页码），属已知限制对照组
    footer_field = section1.footer.add_paragraph("第 ")
    _make_field_paragraph(footer_field, "", instruction=" PAGE ")

    # 页眉内表格
    header_table = section1.header.add_table(rows=1, cols=2, width=4000)
    header_table.cell(0, 0).text = "项目编号"
    header_table.cell(0, 1).text = "CQ-2026-SB-0147"
    header_table2 = section1.header.add_table(rows=1, cols=1, width=4000)
    header_table2.cell(0, 0).text = f"监督电话 {FAKE_PHONE_2}"
    _note("页眉表格", FAKE_PHONE_2)

    # 页眉内 mailto 链接
    header_link = section1.header.add_paragraph()
    add_hyperlink(header_link, f"mailto:{FAKE_EMAIL}", "反馈意见请发邮件")
    _note("页眉超链接", FAKE_EMAIL)

    # 第 2 节：独立页眉页脚
    section2 = document.add_section(WD_SECTION.NEW_PAGE)
    section2.header.is_linked_to_previous = False
    section2.header.paragraphs[0].text = f"附件清单　签署人身份证 {FAKE_ID_VALID}"
    section2.footer.is_linked_to_previous = False
    section2.footer.paragraphs[0].text = "附件无敏感信息"
    document.add_paragraph("附件 A：技术偏离表")
    document.add_paragraph("附件 B：核查记录")
    _note("页眉（第2节）", FAKE_ID_VALID)

    # 第 3 节：保持"链接到前一节"
    document.add_section(WD_SECTION.NEW_PAGE)
    document.add_paragraph("（第 3 节沿用第 2 节页眉页脚，用于验证 part 去重。）")

    # 首页不同
    section1.different_first_page_header_footer = True
    section1.first_page_header.is_linked_to_previous = False
    section1.first_page_header.paragraphs[0].text = f"封面　制表人手机 {FAKE_PHONE_1}"
    section1.first_page_footer.is_linked_to_previous = False
    section1.first_page_footer.paragraphs[0].text = "封面不显示页码"
    _note("首页页眉", FAKE_PHONE_1)

    # 偶数页不同
    section1.even_page_header.is_linked_to_previous = False
    section1.even_page_header.paragraphs[0].text = f"偶数页　对公账号 {FAKE_BANK_CARD}"
    section1.even_page_footer.is_linked_to_previous = False
    section1.even_page_footer.paragraphs[0].text = f"偶数页页脚邮箱 {FAKE_EMAIL}"
    _note("偶数页页眉", FAKE_BANK_CARD)
    _note("偶数页页脚", FAKE_EMAIL)


def _print_manifest(path: Path) -> None:
    """打印"文件 -> 容器 -> 敏感值"对照表，供独立验证逐条核对。"""
    size_kb = path.stat().st_size / 1024
    print(f"已生成：{path}")
    print(f"大小：{size_kb:.1f} KB ({path.stat().st_size} 字节)")
    print()
    print("容器 / 敏感值对照表（数据全部来自 tests/fixtures/make_samples.py 的 FAKE_*）")
    print("-" * 78)

    label_map = {
        "FAKE_ID_VALID": "身份证号",
        "FAKE_BANK_CARD": "银行卡号",
        "FAKE_PHONE_1": "手机号 1(13800138000)",
        "FAKE_PHONE_2": "手机号 2(13912345678)",
        "FAKE_EMAIL": "邮箱",
        "FAKE_IP": "IPv4",
    }
    values = {v: k for k, v in globals().items() if k.startswith("FAKE_")}

    total = 0
    for container, items in EXPECTED_SENSITIVE.items():
        print(f"\n■ {container}")
        for item in items:
            total += 1
            print(f"   - {label_map[values[item]]:<26} {item}")
    print("-" * 78)
    print(f"容器数：{len(EXPECTED_SENSITIVE)}　敏感值条目：{total}")
    print()
    print("已知限制对照组（预期被跳过并计入报告）：")
    print("  - 正文中的 PAGE 域代码段落 1 处")
    print("  - 页脚中的 PAGE 域代码段落 1 处")


if __name__ == "__main__":
    created = build()
    _print_manifest(created)
