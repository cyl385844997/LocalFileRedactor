"""P8 人名识别（jieba）端到端验证。

跑法：
    set PYTHONPATH=src;.verify_libs
    set QT_QPA_PLATFORM=offscreen
    "E:/Program Files/python/python.exe" scripts/e2e_p8.py

场景（6 个）：
    1. ★人名 + 手机号：两者分别正确脱敏（不互吞）
    2. ★大字金额：已支付伍拾万元 / 叁佰伍拾陆万元柒仟捌佰玖拾元 —— 零误伤
    3. 关闭剔除 → 误伤立刻出现（反向对照）
    4. GUI 开关：默认关闭 → 勾选后生效
    5. 真实 docx 端到端
    6. jieba 缺失时的降级（不崩、返回空）
"""
from __future__ import annotations

import os
import sys
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

_REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(_REPO / "src"))
sys.path.insert(0, str(_REPO / ".verify_libs"))

from PySide6 import QtCore, QtWidgets  # noqa: E402

from redactor.core.engine import RedactionEngine  # noqa: E402
from redactor.core.rules import load_rule_dir  # noqa: E402
from redactor.recognizers.name_recognizer import (  # noqa: E402
    build_name_recognizer,
    is_da_zi_amount,
    jieba_available,
)

RULES = load_rule_dir(_REPO / "rules")
PHONE = "13812345678"

DA_ZI_CASES = [
    "已支付伍拾万元整",
    "本次付款金额为叁佰伍拾陆万元柒仟捌佰玖拾元",
]


def _engine(**kw):
    return RedactionEngine(RULES, enable_names=True, **kw)


def _wait(pred, timeout=20.0, step_ms=50):
    import time
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        if pred():
            return True
        QtWidgets.QApplication.processEvents(
            QtCore.QEventLoop.ProcessEventsFlag.AllEvents, step_ms)
    return pred()


def main() -> int:
    print("=" * 74)
    print("P8 人名识别（jieba）端到端验证（v0.8.0）")
    print("=" * 74)
    assert jieba_available(), "jieba 未安装，P8 无法验证"

    # ============================================================
    print("\n--- 场景 1：★人名 + 手机号 分别脱敏 ---")
    text = f"张三13812345678"
    out, matches = _engine().redact(text)
    names = {m.rule_name for m in matches}
    print(f"  in : {text}")
    print(f"  out: {out}")
    print(f"  规则: {sorted(names)}")
    assert PHONE not in out, f"手机号被吞：{out}"
    assert "[手机号]" in out
    assert "张三" not in out, f"人名没脱：{out}"
    assert "[人名]" in out
    assert names == {"手机号", "人名"}, names
    print("  ✅ 场景 1 通过：两者分别脱敏，互不吞并")

    text2 = "项目负责人是李四，电话13812345678，公司在北京"
    out2, names2 = _engine().redact(text2)
    print(f"\n  组合场景")
    print(f"  in : {text2}")
    print(f"  out: {out2}")
    print(f"  规则: {sorted({m.rule_name for m in names2})}")
    assert "[人名]" in out2 and "[手机号]" in out2
    assert "李四" not in out2 and PHONE not in out2
    print("  ✅ 组合场景通过")

    # ============================================================
    print("\n--- 场景 2：★大字金额 零误伤 ---")
    for t in DA_ZI_CASES:
        out, _ = _engine().redact(t)
        print(f"  {t}")
        print(f"    -> {out!r}")
        assert "[人名]" not in out, f"★大字金额被误脱：{out}"

    # 顺带把探针里那 9 条干净负例也跑一遍
    for t in ["甲方为重庆创锐电气技术有限公司",
              "我方与西南政法大学签署协议",
              "本项目由重庆水务集团实施",
              "该项目建设周期为十八个月",
              "合同总价为壹佰贰拾万元",
              "第五条第二款约定",
              "根据民法典第九百九十条",
              "这个项目很难做",
              "他今天去公司了"]:
        out, _ = _engine().redact(t)
        assert "[人名]" not in out, f"{t!r} 误伤：{out!r}"
    print(f"  ✅ 场景 2 通过：大字金额 + 9 条负例，零误伤")

    # ============================================================
    print("\n--- 场景 3：★反向对照（关闭剔除 → 误伤出现）---")
    evil = build_name_recognizer(strip_da_zi=False)
    evil_hits = [m.raw for t in DA_ZI_CASES for m in evil.find(t)]
    good_hits = [m.raw for t in DA_ZI_CASES
                 for m in build_name_recognizer(strip_da_zi=True).find(t)]
    print(f"  剔除开启命中: {good_hits}")
    print(f"  剔除关闭命中: {evil_hits}")
    assert good_hits == [], f"真版不该有命中，却有 {good_hits}"
    assert evil_hits, "❌ 对照失效：关掉剔除后仍零误伤"
    assert set(evil_hits) == {"伍拾", "陆"}, evil_hits
    print("  ✅ 场景 3 通过：剔除逻辑确实在起作用")

    # ============================================================
    print("\n--- 场景 4：GUI 开关（默认关闭）---")
    app = QtWidgets.QApplication.instance() or QtWidgets.QApplication(sys.argv)
    from redactor.gui import MainWindow

    w = MainWindow()
    print(f"  人名识别开关默认勾选: {w.names_action().isChecked()}")
    print(f"  大字剔除开关默认勾选: {w.strip_da_zi_action().isChecked()}")
    print(f"  大字剔除开关默认可用: {w.strip_da_zi_action().isEnabled()}")
    assert w.names_action().isChecked() is False, "★必须默认关闭"
    assert w.strip_da_zi_action().isEnabled() is False

    w.names_action().setChecked(True)
    app.processEvents()
    print(f"  勾选后 controller: {w.controller().names_enabled()}")
    print(f"  勾选后大字开关可用: {w.strip_da_zi_action().isEnabled()}")
    print(f"  日志: {w.log_panel().toPlainText().strip()}")
    assert w.controller().names_enabled() is True
    assert "人名识别已启用" in w.log_panel().toPlainText()
    w.close()
    print("  ✅ 场景 4 通过")

    # ============================================================
    print("\n--- 场景 5：真实 docx 端到端 ---")
    import tempfile
    from docx import Document

    tmp = Path(tempfile.mkdtemp(prefix="redactor_p8_"))
    src = tmp / "in" / "contract.docx"
    src.parent.mkdir(parents=True, exist_ok=True)
    d = Document()
    d.add_paragraph("甲方代表张三，身份证110101199001011237")
    d.add_paragraph("已支付伍拾万元整")
    d.add_paragraph("联系电话13812345678")
    d.add_paragraph("乙方为重庆创锐电气技术有限公司")
    d.save(str(src))

    w2 = MainWindow()
    w2.drop_zone().add_files([str(src)])
    w2.names_action().setChecked(True)
    app.processEvents()
    out_root = tmp / "out"
    assert w2.start_processing(output_root=out_root)
    assert _wait(lambda: w2.state() == "done")
    w2.close()

    result = out_root / "contract_redacted.docx"
    assert result.exists(), f"输出文件不存在：{result}"
    text_out = "\n".join(p.text for p in Document(str(result)).paragraphs)
    print("  --- 脱敏后正文 ---")
    for line in text_out.splitlines():
        print(f"    {line}")
    print("  -------------------")
    assert "张三" not in text_out, "人名没脱"
    assert "[人名]" in text_out
    assert "伍拾" in text_out, "★大字金额被误脱（'已支付伍拾万元整'该原样保留）"
    assert "13812345678" not in text_out, "手机号没脱"
    assert "110101199001011237" not in text_out, "身份证没脱"
    # 注意：本场景**没开范围脱敏**，公司名不该被脱 —— 它只是 jieba 的
    # 一个真实误伤源？实测"重庆创锐电气技术有限公司"里没有 nr 词，
    # 所以人名识别不会碰它（这也正是范围脱敏 P7 存在的原因）。
    assert "重庆创锐电气技术有限公司" in text_out, \
        "没开范围脱敏时公司名应原样保留（它是 P7 的职责，不是 P8）"
    print("  ✅ 场景 5 通过：人名脱敏、大字金额保留、格式规则正常")
    print("     （公司名保留属预期：范围脱敏未开启）")

    # ============================================================
    print("\n--- 场景 6：jieba 缺失时降级 ---")
    import builtins
    real_import = builtins.__import__

    def fake_import(name, *a, **k):
        if name.startswith("jieba"):
            raise ImportError("模拟未安装")
        return real_import(name, *a, **k)

    builtins.__import__ = fake_import
    try:
        from redactor.recognizers.name_recognizer import NameRecognizer
        assert build_name_recognizer() is None, "缺依赖时应返回 None"
        assert jieba_available() is False
        assert NameRecognizer().find("张三来了") == [], "缺依赖时应返回空而非崩"
        print("  ✅ 场景 6 通过：返回 None / 空，不抛异常")
    finally:
        builtins.__import__ = real_import

    import shutil
    shutil.rmtree(tmp, ignore_errors=True)

    print("\n" + "=" * 74)
    print("✅ P8 端到端验证全部通过")
    print("=" * 74)
    return 0


if __name__ == "__main__":
    sys.exit(main())