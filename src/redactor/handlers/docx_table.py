"""docx 表格遍历（优先级 3）。

为什么需要独立模块
------------------
``document.paragraphs`` **不包含表格内的段落**（实测确认），所以表格内容
必须单独遍历，否则整张表的内容都会漏脱敏。

三个实测确认的边界情况
---------------------
1. **合并单元格**::

       cell(0,0) is cell(0,1)  ->  True

   合并后的单元格是**同一个 XML 元素**，按坐标遍历会访问它多次。
   重复处理不会破坏内容（第二次无命中），但会**虚增命中次数**，
   报告里就会多记一次。因此必须按 XML 元素去重。

2. **嵌套表格**：嵌套表**不在** ``document.tables`` 里，
   也不在外层单元格的 ``paragraphs`` 里，必须递归 ``cell.tables``。

3. **单元格内多段落**：一个单元格可以有多个段落，逐个处理。

跳过策略
--------
**只跳过问题段落，不整表跳过。** 表格常有几十行，域代码通常只在个别单元格
（如页码栏）；整表跳过会连带漏掉同表其他单元格的敏感信息。
"""

from __future__ import annotations

from typing import Dict, Iterator, List, Optional, Set, Tuple

from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

#: 嵌套表格递归深度上限。防止病态文档造成无限递归或超长位置标签。
#: 实测 Word 自身嵌套也很少超过 3 层，留 5 层余量足够。
MAX_NESTED_DEPTH = 5


def cell_position_label(table_index: int, row_index: int, col_index: int, prefix: str = "") -> str:
    """生成单元格位置标签，如 ``表1 R2C3`` 或 ``表1 R2C3 内表1 R1C1``。

    Args:
        table_index: 表格序号（从 1 开始，同一层内计数）。
        row_index: 行号（从 1 开始）。
        col_index: 列号（从 1 开始）。
        prefix: 外层位置前缀（嵌套表时传入）。

    Returns:
        紧凑的位置标签。用R/C 缩写而非"第X行第Y列"，
        因为几百个单元格时中文全称会让报告臃肿。
    """
    own = f"表{table_index} R{row_index}C{col_index}"
    return f"{prefix} 内{own}" if prefix else own


def iter_table_paragraphs(
    table: Table,
    table_index: int,
    *,
    prefix: str = "",
    depth: int = 1,
) -> Iterator[Tuple[Paragraph, str]]:
    """遍历表格内所有段落，产出 ``(段落, 位置标签)``。

    会递归处理嵌套表格，并**按 XML 元素去重**以应对合并单元格。

    Args:
        table: 目标表格。
        table_index: 本表格在同层中的序号（从 1 开始）。
        prefix: 外层位置前缀（嵌套表时传入）。
        depth: 当前递归深度（从 1 开始）。

    Yields:
        ``(段落对象, 位置标签)``。位置标签形如 ``表1 R2C3``。

    Raises:
        RecursionError: 嵌套深度超过 :data:`MAX_NESTED_DEPTH`。
    """
    if depth > MAX_NESTED_DEPTH:
        raise RecursionError(
            f"嵌套表格深度超过上限 {MAX_NESTED_DEPTH}，疑似病态文档，已停止处理该表"
        )

    # 记录已处理的单元格 XML 元素，避免合并单元格被重复处理。
    #
    # **必须用元素本身，不能用 id()**。
    # 实测教训：python-docx 的 ``cell._tc`` 是代理属性，遍历过程中旧代理会被
    # 垃圾回收，其 id() 随后可能被**新对象复用**——导致后面的正常单元格
    # 被误判为"已访问"而整段漏处理（实测丢了整个 R2C3）。
    # lxml 元素本身可哈希且按身份比较，直接放进 set 才是正确做法。
    seen_cells: Set[object] = set()

    for row_index, row in enumerate(table.rows, start=1):
        for col_index, cell in enumerate(row.cells, start=1):
            key = cell._tc
            if key in seen_cells:
                continue  # 合并单元格：同一 XML 元素已处理过
            seen_cells.add(key)

            label = cell_position_label(table_index, row_index, col_index, prefix)

            # 先处理本单元格的段落（不含嵌套表内容）
            for paragraph in cell.paragraphs:
                yield paragraph, label

            # 再递归处理单元格内的嵌套表格
            for nested_index, nested_table in enumerate(cell.tables, start=1):
                yield from iter_table_paragraphs(
                    nested_table,
                    nested_index,
                    prefix=label,
                    depth=depth + 1,
                )


def iter_all_table_paragraphs(document) -> Iterator[Tuple[Paragraph, str]]:
    """遍历文档里所有表格（含嵌套）内的段落。

    Args:
        document: python-docx 的 ``Document`` 对象。

    Yields:
        ``(段落对象, 位置标签)``。
    """
    for table_index, table in enumerate(document.tables, start=1):
        yield from iter_table_paragraphs(table, table_index)